import fitz  # สำหรับอ่าน PDF
import whisper
from langchain_ollama import OllamaLLM
from langchain_core.prompts import PromptTemplate
from langchain.chains import LLMChain
from docx import Document
import tempfile
import requests
from transformers import WhisperProcessor, WhisperForConditionalGeneration
from moviepy.editor import VideoFileClip
import librosa
import torchaudio
import torch
import noisereduce as nr
import math


# อ่านข้อความจาก PDF ภาษาไทย
def extract_text_from_pdf(pdf_path):
    doc = fitz.open(pdf_path)
    return "\n".join(page.get_text() for page in doc)


# ถอดเสียงจากวิดีโอ URL ด้วย HuggingFace Whisper


def transcribe_video_url_with_transformers(url, model_name="scb10x/monsoon-whisper-medium-gigaspeech2", chunk_duration=15):
    video_path = download_video_from_url(url)

    # แปลงวิดีโอเป็นเสียง .wav
    audio_path = tempfile.NamedTemporaryFile(delete=False, suffix=".wav").name
    video = VideoFileClip(video_path)
    video.audio.write_audiofile(audio_path, logger=None)

    waveform, sample_rate = torchaudio.load(audio_path)
    if sample_rate != 16000:
        waveform = torchaudio.transforms.Resample(orig_freq=sample_rate, new_freq=16000)(waveform)

    # ลดเสียงรบกวน (ใช้ช่วงแรกเป็น noise profile)
    noise_sample = waveform[:, :int(sample_rate * 0.5)]
    reduced_waveform = nr.reduce_noise(y=waveform.numpy()[0], y_noise=noise_sample.numpy()[0], sr=sample_rate)

    # เตรียม model
    device = "cuda" if torch.cuda.is_available() else "cpu"
    processor = WhisperProcessor.from_pretrained(model_name)
    model = WhisperForConditionalGeneration.from_pretrained(model_name)
    model.to(device)
    model.eval()

    # ตัด waveform เป็น chunk แล้วประมวลผลทีละช่วง
    total_samples = len(reduced_waveform)
    chunk_size = chunk_duration * sample_rate
    num_chunks = math.ceil(total_samples / chunk_size)

    all_text = []

    for i in range(num_chunks):
        start = i * chunk_size
        end = min(start + chunk_size, total_samples)
        chunk = reduced_waveform[start:end]

        input_tensor = torch.tensor(chunk, dtype=torch.float32).unsqueeze(0).to(device)
        inputs = processor(input_tensor.squeeze().cpu().numpy(), sampling_rate=16000, return_tensors="pt").input_features.to(device)

        with torch.no_grad():
            predicted_ids = model.generate(inputs)
            text = processor.batch_decode(predicted_ids, skip_special_tokens=True)[0]
            all_text.append(text)
        print(text)

    return "\n".join(all_text)



# ดาวน์โหลดวิดีโอจาก URL

def download_video_from_url(url):
    response = requests.get(url, stream=True)
    response.raise_for_status()

    tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".mp4")
    for chunk in response.iter_content(chunk_size=8192):
        if chunk:
            tmp_file.write(chunk)
    tmp_file.close()
    return tmp_file.name


# ถอดเสียงจากวิดีโอ URL ด้วย OpenAI Whisper

def transcribe_video_url(url):
    video_path = download_video_from_url(url)
    model = whisper.load_model("large", device='cuda')
    result = model.transcribe(video_path, language="th")
    return result["text"]


# สรุปข้อความโดยใช้ LLM

def summarize_text(text):
    llm = OllamaLLM(model="scb10x/llama3.1-typhoon2-8b-instruct:latest")
    prompt = PromptTemplate.from_template(
        """คุณคือระบบช่วยสร้างข้อมูลสำหรับจัดเก็บในเวกเตอร์ดาต้าเบส (Vector Database) เช่น FAISS หรือ Weaviate

จากเนื้อหาที่ได้รับต่อไปนี้ หากข้อความดูไม่เหมาะสมหรือไม่น่าเป็นไปได้สามารถแก้ข้อความได้เองเลยย เนื้อหานี้เกี่ยวกับ Archive e-Mail : แยก e-Mail รายปี (Outlook2019):

{content}

ให้คุณสร้างข้อมูลใหม่ในรูปแบบ JSON ที่ประกอบด้วย:

{{
  "title": "ชื่อเรื่องหรือหัวข้อหลักที่สื่อถึงเนื้อหา",
  "summary": "สรุปเนื้อหาทั้งหมดให้กระชับและเข้าใจง่าย",
  "keywords": "keyword1", "keyword2", "keyword3", ...,  // คำค้นสำคัญทั้งเชิงเนื้อหาและเทคนิค
  "category": "หมวดหมู่ของเนื้อหา IT, excel, AXุ4, Outlook"
}}

หากข้อความบางช่วงไม่สมบูรณ์หรือผิด ให้ปรับให้อ่านรู้เรื่องและตรงประเด็น
ห้ามเขียนคำอธิบายหรือเกริ่นนำใดๆ ให้คืนค่าเฉพาะ JSON เท่านั้น"""
    )
    chain = LLMChain(llm=llm, prompt=prompt)
    return chain.run(content=text)


# เขียนข้อความลง Word

def write_to_word(text, filename):
    doc = Document()
    doc.add_heading("สรุปรายงาน", 0)
    doc.add_paragraph(text)
    doc.save(filename)


# 🔧 ตัวอย่างการใช้งานจริง

# url = "https://ibot-ncr.dyndns.org/Archive-Email-Outlook2019.mp4"

# # ใช้ transformers (scb10x)
# thai_text = transcribe_video_url_with_transformers(url)

# # หรือใช้ Whisper OpenAI
# # thai_text = transcribe_video_url(url)

# print("\n📄 ถอดเสียงสำเร็จ:\n", thai_text)


# thai_text = """ สวัสดีครับพีเดอวันนี้เรามีเรื่องของการแอบอีเมลรายปีของไมโครซอฟเอาลุค๒๐๑๙มาแนะนํานะครับเริ่มต้นให้เราเปิดเอาลุคขึ้นมา
# ไปที่เมนูไฟล์ครับไปที่เมนูทูครับเลือกทีมอัพโอไอเทมครับให้เราเลือกที่อาคาริดิสโคเดอร์แอนอร์ซับโคเดอร์นะครับเลือกที่อีเมลปัจจุบันของเราครับแล้วก็มาเลือกวันที่นะคร
# ับวันที่ให้เลือกเป็นวันที่๑เดือนหนึ่งนะครับปีปัจจุบันครับในส่วนตรงนี้เราจะเลือกนะครับว่าเราจะเก็บอาคารโปรเดอร์ไว้ที่ไหนผมจะเก็บไว้ที่ได้ดีนะครับโปรเดอร์อีเมลอลุ
# กนะครับแล้วก็เปลี่ยนชื่อไฟล์เป็นปีก่อนหน้านี้นะครับ
# หลังจากนั้นกดโอเคครับแล้วก็กดโอเคอีกหนึ่งครับเมื่อเรากลับมาที่หน้าอาลุปปกตินะครับเราจะเห็นว่ามีอาคารโปรเศษมันลันอยู่นะครับแต่ระหว่างนี้เราสามารถใช้งานอีเมลได้ตา
# มปกติครับเมื่ออาคารโปรเศษที่ลันอยู่มันหายไปนะครับแต่ว่าตอนนี้มันทํางานอาคารอีเมลให้เราเรียบร้อยแล้วนะครับให้เราไปที่เมนูไฟล์ครับเลือกแอคเคานเซตติ้งแล้วก็เลือกแอ
# คเคานเซติ้งอีกครั้งนี้ไปที่แถบต่างต้าไฟล์ครับไปเลือก
# แล้วก็ไปเลือกเซตติ้งครับแล้วก็เปลี่ยนตรงช่องเนมตรงนี้นะครับให้เป็นอีเมลปีก่อนหน้านี้นะครับในที่นี้ผมเปลี่ยนไว้แล้วนะครับหลังจากนั้นกดโอเคครับแล้วก็กดโค้ชไม่ได้เ
# ลยครับในส่วนนี้เสร็จในเรื่องของการอักขายอีเมลเรียบร้อยแล้วนะครับต่อไปนะครับเราจะเข้าไปทําการคองแพคอีเมลนะครับเพื่อเป็นการลดขนาดอีเมลไฟล์ของเรานะครับโดยที่เข้าไป
# ที่เมนูไฟล์ครับได้แอคาซิตติ้งแล้วก็แอคาซิตเนสอีกครั้งหนึ่งเพื่อที่ได้ทั่งไปครับแล้วก็
# เลือกอีเมลปัจจุบันของเรานะครับสังเกตง่ายอีเมลปัจจุบันของเราจะมีติดถูกนะครับให้เราไฮไลท์แล้วก็กดเซตติ้งนะครับแล้วก็ไปเลือกที่คอมแคนแล้วนะครับในส่วนของการคอมแพคนี
# ้นะครับบางเครื่องอาจจะใช้เวลาค่อนข้างนานนะครับระหว่างที่คอมแพคอยู่ก็ไม่สามารถใช้งานอีเมลนะครับหรือถ้าให้กดคอมแพคไปแล้วนะครับแล้วต้องการใช้งานอีเมลนะครับให้เราเซ
# ็นเซินก่อนได้แล้วก็กลับมาโค้แท็กใหม่อีกรอบหนึ่งก็ได้เพื่อดําเนินการเรียบร้อยแล้วกดโอเค
# อย่าลืมติดตามพี่ทูเอในคลิปถัดไปนะครับสําหรับวันนี้สวัสดีครับ"""

# summary = summarize_text(thai_text)
# print(summary)
# # write_to_word(summary, "รายงาน.docx")
# # print("\n✅ สรุปเสร็จสิ้น บันทึกไฟล์แล้ว: รายงาน.docx")
